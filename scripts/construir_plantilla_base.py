# -*- coding: utf-8 -*-
"""Construye assets/plantilla_base.xlsx: memo real "en blanco" con TOKENS sembrados.

Estrategia (robusta, no depende de que ninguna libreria JS entienda el membrete):
la plantilla conserva TODO el formato/logos de un memo real, pero en cada celda
rellenable se siembra un token unico tipo @@CT_L1_ID@@. En el navegador, la app
descomprime el .xlsx (es un zip), reemplaza los tokens como TEXTO en el XML y
vuelve a comprimir -- sin tocar xl/media ni xl/drawings, asi los logos quedan
intactos byte a byte.

openpyxl al escribir tokens o borrar hojas preserva las imagenes de las hojas
que quedan (verificado); solo DUPLICAR hojas las perderia, y aca no se duplica.

Uso:
    python construir_plantilla_base.py [memo_origen.xlsx] [salida.xlsx]
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import MergedCell

RAIZ = Path(__file__).resolve().parents[1]
ORIGEN_DEFAULT = (
    r"C:\Users\carlos.venegas\OneDrive - Sernageomin\05 Central Los Cipreses"
    r"\03 Datos y Resultados\01 Laboratorio\01 Solicitudes e ingresos\Memos\2026"
    r"\Memo N°88-89-90 muestras.xlsx"
)
SALIDA_DEFAULT = RAIZ / "assets" / "plantilla_base.xlsx"

# --- FichaIngreso: columnas de datos y rango de filas a sembrar ---
FICHA_INGRESO_FILA_INI = 11
FICHA_INGRESO_FILAS = 24                       # filas 11..34
FICHA_INGRESO_COLS = {
    "item": "B", "requerimiento": "C", "id": "E", "proyecto": "F",
    "centro": "G", "colector": "H", "este": "I", "norte": "J",
    "tipomuestra": "K", "litologia": "L", "unidad": "M", "fecha": "N", "obs": "O",
}

# --- Hojas Memo por tipo: codigo corto por hoja para tokens unicos ---
# (el nombre de hoja lleva "N°" con caracter grado; se localiza por substring)
MEMO_SHEETS = {"CT": "CT", "U-Pb": "UPB", "RX": "RX"}
MEMO_CELDA_FECHA = "D14"
MEMO_CELDA_FELAB = "G4"
MEMO_LISTADO_FILA_INI = 21
MEMO_LISTADO_FILAS = 6                          # filas 21..26 antes de la firma
MEMO_LISTADO_PARES = [("C", "D"), ("E", "F")]   # hasta 12 muestras por memo

# --- Ficha Datacion (pool): celdas a sembrar por hoja ---
FICHA_DAT_CELDAS = {
    "solicitante": "C10", "id": "C11", "litologia": "C14",
    "unidad": "C16", "clasificacion": "C25",
}
# celdas de juicio manual del geologo: se dejan en blanco (no token)
FICHA_DAT_BLANQUEAR = ["C17", "C18", "C28", "C29", "B31", "B15", "B18", "C20", "E20"]


def _set(ws, ref, valor):
    if isinstance(ws[ref], MergedCell):
        return False
    ws[ref] = valor
    return True


def sembrar_ficha_ingreso(ws):
    for i in range(1, FICHA_INGRESO_FILAS + 1):
        fila = FICHA_INGRESO_FILA_INI + i - 1
        for campo, col in FICHA_INGRESO_COLS.items():
            _set(ws, f"{col}{fila}", f"@@FI_{campo}_{i}@@")


def sembrar_memo_tipo(ws, code):
    _set(ws, MEMO_CELDA_FECHA, f"@@{code}_FECHA@@")
    _set(ws, MEMO_CELDA_FELAB, f"@@{code}_FELAB@@")
    # G8 (numero de memorando) se deja manual: la plantilla mantiene "Memorando Nº "
    n = 0
    for (cnum, cid) in MEMO_LISTADO_PARES:
        for r in range(MEMO_LISTADO_FILA_INI, MEMO_LISTADO_FILA_INI + MEMO_LISTADO_FILAS):
            n += 1
            _set(ws, f"{cnum}{r}", f"@@{code}_L{n}_N@@")
            _set(ws, f"{cid}{r}", f"@@{code}_L{n}_ID@@")


def sembrar_ficha_datacion(ws, idx):
    for campo, ref in FICHA_DAT_CELDAS.items():
        _set(ws, ref, f"@@FD{idx}_{campo}@@")
    for ref in FICHA_DAT_BLANQUEAR:
        _set(ws, ref, None)


def construir(origen, salida):
    wb = openpyxl.load_workbook(origen)
    n_ficha = 0
    for nombre in list(wb.sheetnames):
        ws = wb[nombre]
        limpio = nombre.strip()
        if limpio == "FichaIngreso":
            sembrar_ficha_ingreso(ws)
        elif limpio.startswith("Memo N"):
            code = next((c for k, c in MEMO_SHEETS.items() if k in limpio), None)
            if code:
                sembrar_memo_tipo(ws, code)
        elif limpio.startswith("Ficha Dataci"):
            n_ficha += 1
            sembrar_ficha_datacion(ws, n_ficha)
            ws.title = f"Ficha Datación {n_ficha}"
        # DatosProyecto: se deja tal cual (valores por defecto del proyecto)

    salida = Path(salida)
    salida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(salida)

    wb2 = openpyxl.load_workbook(salida)
    print(f"Plantilla base (con tokens) guardada: {salida}")
    print(f"Hojas ({len(wb2.sheetnames)}):")
    for n in wb2.sheetnames:
        ws = wb2[n]
        print(f"  - {n!r}  imgs={len(getattr(ws, '_images', []))}")
    print(f"Pool de fichas de datacion: {n_ficha}")


if __name__ == "__main__":
    origen = sys.argv[1] if len(sys.argv) > 1 else ORIGEN_DEFAULT
    salida = sys.argv[2] if len(sys.argv) > 2 else SALIDA_DEFAULT
    construir(origen, salida)
